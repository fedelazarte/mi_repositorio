"""CV orientado a una oferta concreta, armado solo con datos del perfil.

No inventa empleos, fechas ni habilidades. Reordena lo que ya está en `perfil.yaml`
para que lo que la oferta pide quede adelante. Con `--llm` un modelo puede reescribir
la redacción, con la misma restricción.
"""
from __future__ import annotations

import re
from pathlib import Path

from . import config
from .matcher import canon, extract_required_skills, normalize
from .models import Job
from .profile import Profile

CVS_DIR = config.HOME / "cvs"


def _slug(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", normalize(text)).strip("-")
    return slug[:40] or "rol"


def _relevant_skills(profile: Profile, job: Job) -> tuple[list[str], list[str], list[str]]:
    """Habilidades del perfil que la oferta pide, el resto, y las que estás aprendiendo y la oferta pide."""
    required = extract_required_skills(normalize(job.full_text), profile)
    relevant, other = [], []
    for skill in profile.habilidades:
        (relevant if canon(skill) in required else other).append(skill)
    learning = [skill for skill in profile.aprendiendo if canon(skill) in required]
    return relevant, other, learning


def render_cv(profile: Profile, job: Job) -> str:
    relevant, other, learning = _relevant_skills(profile, job)
    contact = " · ".join(
        part for part in (profile.email, profile.telefono, profile.linkedin, profile.ubicacion_actual) if part
    )
    lines = [f"# {profile.nombre or 'CV'}", "", profile.titulo_actual, ""]
    if contact:
        lines.extend([contact, ""])
    lines.extend([
        f"Preparado para: {job.title} — {job.company}",
        "",
        "## Perfil",
        "",
        _summary(profile, job, relevant),
        "",
        "## Habilidades",
        "",
    ])
    if relevant:
        lines.append("Para este puesto: " + ", ".join(relevant))
        lines.append("")
    if other:
        lines.append("También: " + ", ".join(other))
        lines.append("")
    if learning:
        lines.append("En formación, y este puesto lo pide: " + ", ".join(learning))
        lines.append("")
    lines.extend(["## Experiencia", ""])
    if not profile.experiencia:
        lines.append("Sin experiencia cargada en el perfil.")
        lines.append("")
    for role in profile.experiencia:
        puesto = role.get("puesto") or ""
        empresa = role.get("empresa") or ""
        periodo = role.get("periodo") or ""
        header = " — ".join(part for part in (puesto, empresa) if part)
        if periodo:
            header = f"{header} ({periodo})" if header else periodo
        lines.append(f"### {header}")
        lines.append("")
        description = (role.get("descripcion") or "").strip()
        if description:
            lines.extend([description, ""])
    education = [str(item) for item in (profile.raw.get("educacion") or []) if str(item).strip()]
    if education:
        lines.extend(["## Educación", ""])
        lines.extend(f"- {item}" for item in education)
        lines.append("")
    if profile.idiomas:
        lines.extend(["## Idiomas", "", ", ".join(profile.idiomas), ""])
    return "\n".join(lines).rstrip() + "\n"


def _summary(profile: Profile, job: Job, relevant: list[str]) -> str:
    base = " ".join(profile.resumen.split())
    opening = f"{profile.titulo_actual} con {profile.anios_experiencia} años de experiencia." if profile.anios_experiencia else profile.titulo_actual
    highlight = f" Para {job.title} en {job.company} destaco: {', '.join(relevant)}." if relevant else ""
    text = " ".join(part for part in (opening, base) if part).strip()
    return (text + highlight).strip()


def rewrite_with_llm(profile: Profile, job: Job, draft: str) -> str:
    """Pide una redacción más dirigida. Si falla o no hay clave, devuelve el borrador."""
    import json

    import requests

    if not config.OPENAI_API_KEY:
        return draft
    prompt = f"""Reescribí este CV en markdown para la oferta de {job.title} en {job.company}.
Reglas:
- Usá exclusivamente hechos que ya están en el CV. No inventes empleos, fechas, herramientas, métricas ni estudios.
- No cambies los nombres de empresas ni los períodos.
- El resumen tiene que hablarle a este puesto.
- Devolvé solo el markdown, empezando con "# ".

OFERTA:
{job.title} — {job.company}
{(job.description or "")[:4000]}

CV:
{draft}
"""
    try:
        response = requests.post(
            f"{config.OPENAI_BASE_URL}/chat/completions",
            headers={"Authorization": f"Bearer {config.OPENAI_API_KEY}"},
            json={
                "model": config.OPENAI_MODEL,
                "temperature": 0.3,
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=60,
        )
        response.raise_for_status()
        text = response.json()["choices"][0]["message"]["content"].strip()
        text = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", text)
        if not text.startswith("# "):
            return draft
        return text if text.endswith("\n") else text + "\n"
    except (requests.RequestException, KeyError, IndexError, json.JSONDecodeError):
        return draft


def markdown_to_docx(markdown: str, path: Path) -> None:
    from docx import Document

    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=0)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip():
            document.add_paragraph(line.strip())
    document.save(path)


def write_cv(profile: Profile, job: Job, directory: Path | None = None, *, use_llm: bool = False) -> tuple[Path, Path]:
    folder = directory or CVS_DIR
    folder.mkdir(parents=True, exist_ok=True)
    markdown = render_cv(profile, job)
    if use_llm:
        markdown = rewrite_with_llm(profile, job, markdown)
    stem = f"{job.id}-{_slug(job.title)}-{_slug(job.company)}"
    md_path = folder / f"{stem}.md"
    docx_path = folder / f"{stem}.docx"
    md_path.write_text(markdown, encoding="utf-8")
    markdown_to_docx(markdown, docx_path)
    return md_path, docx_path
